#!/usr/bin/env python3
"""Find likely duplicate ParishSoft Families in an Archdiocese-wide export.

One-off analysis tool (not part of ParishKit).  Reads the Family-level and
Member-level export XLSXs (plus, optionally, the "enhanced" Family export
for extra fields), finds candidate duplicate pairs with explicit,
explainable rules, groups them, picks a surviving DUID per group, and
writes a report workbook plus a short Markdown summary.

Usage:
  dedup.py --families FAMILIES.xlsx --members MEMBERS.xlsx [--enhanced ENHANCED.xlsx] --out-dir DIR

See README.md for the full rules and CLAUDE.md for maintainer notes.

Evidence comes from two levels:
- Family fields: names, address, phones (incl. emergency phone), emails.
- Members: the same person (name + exact birthdate) in both Families is the
  strongest evidence.  Member roles separate true duplicates from related
  Families that legitimately share a person:
    * each Family has a living partner the other lacks -> divorce /
      remarriage (or different households): related, not a duplicate;
    * the shared person is an ex-spouse in one Family -> related;
    * the shared person is a partner (head/spouse) in one Family but a
      child/other in the other -> grown child, or parent who moved in:
      related;
    * same-named partners with clearly different birthdates (Jr./Sr.) ->
      different people: vetoes a Family-level match.
"""

import argparse
import collections
import datetime
import heapq
import os
import pickle
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from rapidfuzz.distance import JaroWinkler, Levenshtein

# ---------------------------------------------------------------------------
# Tunables

# A contact value (email, phone, address) shared by more than this many
# distinct surnames is treated as shared (parish office phone, nursing home,
# apartment building) and is not identity evidence on its own.
MAX_SURNAMES_PER_CONTACT = 3
# Skip candidate blocks larger than this (pairs grow quadratically).
MAX_BLOCK = 300
# A calendar day on which at least this many Families were unregistered or
# "last updated" is a bulk job, not activity by those Families.
BULK_DAY_MIN = 500
# ... or at least this many Families in one parish on one day (a parish-wide
# bulk job too small to cross the Archdiocese-wide threshold).
BULK_PARISH_DAY_MIN = 100
# A minute in which at least this many Members were modified is a bulk
# update (data conversion, mass cleanup), not Family activity.
BULK_MODIFIED_PER_MINUTE = 200
# Groups larger than this are probably chained through weak links.
MAX_GROUP = 8
# Jaro-Winkler similarity for "similar last name" (typos).
SIMILAR_LAST = 0.92
# A Family whose members carry at least this many distinct surnames is a
# catch-all record (e.g. "Certificate People"), so its members are not
# evidence about anyone's household.
CATCHALL_SURNAMES = 6
# An "Unknown"/"Adult"-typed member this many years younger than another
# member of the same Family is treated as a child of that Family.
GENERATION_YEARS = 18
TODAY = datetime.date.today()

CONFIDENCE = {1: "Very high", 2: "High", 3: "Medium", 4: "Low (review)"}

COLS = {
    "duid": "Families Family DUID",
    "env": "Env #",
    "registered": "Is Registered",
    "mailing": "Families Mailing Name",
    "first": "Families First Names",
    "last": "Families Last Name",
    "group": "Group Name",
    "home_phone": "Home Phone",
    "addr1": "Primary Address 1",
    "addr2": "Primary Address 2",
    "city": "Primary City",
    "primary_phone": "Primary Phone",
    "zip": "Primary Postal Code (incl_ +4)",
    "state": "Contact Info Primary State",
    "email": "Email Address",
    "org_id": "Organization ID",
    "org": "Organization Name",
    "unregistered": "Date Unregistered",
    "membership": "Families Membership Date",
}

MEMBER_COLS = {
    "fam": "Families Family DUID",
    "duid": "Members Member DUID",
    "cell": "Members_Contact Info Cell Phone",
    "email": "Members_Contact Info Email Address",
    "birth": "Members Birthdate",
    "death": "Members Date of Death",
    "modified": "Members Date Modified",
    "first": "Members First Name (Formal)",
    "nick": "Members First Name (Informal/Nickname)",
    "middle": "Members Middle Name",
    "status": "Members Member Status",
    "deceased": "Members is Deceased",
    "last": "Members Last Name",
    "maiden": "Members Maiden Name",
    "type": "Members Member Type",
    "suffix": "Members Suffix",
}

ENHANCED_COLS = {
    "duid": "Families Family DUID",
    "registered": "Families is Registered",
    "last_updated": "Families Date Last Updated",
    "emergency": "Families_Contact Info Emergency Phone",
    "emergency_desc": "Families_Contact Info Emergency Phone Desc",
}
# Emergency phones described as belonging to someone outside the household
# are not evidence about the Family's identity.
NOT_OWN_PHONE = re.compile(
    r"grand|aunt|uncle|neighbo|friend|sister|brother|sibling|cousin|in.law|nana|papa|"
    r"parent|daughter|son\b|niece|nephew|babysit|pastor|office|church|school",
    re.I,
)

# Member Type classification.  "partner" = an adult who heads the household
# (head/spouse); "child" = a child of the household; "ex" = a former spouse;
# "other" = any other relative or adult (mother, sibling, ...).  Unknown /
# Adult types are resolved per Family by age (see assign_roles).
PARTNER_TYPES = {
    "head",
    "husband",
    "wife",
    "spouse",
    "spouse/significant other",
    "significant other",
    "partner",
    "fiance",
    "fiancee",
    "self",
    "head of household",
    "dec. husband",
    "dec. wife",
    "widow",
    "widower",
    "girlfriend",
    "boyfriend",
}
# Resolved by age within the Family (see assign_roles).  Mother/Father stay
# "other": in a Family they are usually a parent living with an adult child.
AMBIGUOUS_TYPES = {
    "unknown",
    "adult",
    "other",
    "",
    "non-catholic",
    "sing. ad",
    "single adult",
    "mom",
    "dad",
    "m",
    "f",
}
# Truncated or variant child types ("grandchi", "adult daug", "step-son").
CHILD_PREFIXES = (
    "child",
    "son",
    "daug",
    "grandc",
    "grands",
    "grandd",
    "step",
    "adult d",
    "adult s",
    "adult c",
    "foster",
)
CHILD_TYPES = {
    "child",
    "son",
    "daughter",
    "grandchild",
    "granddaughter",
    "grandson",
    "adopted",
    "young adult",
    "adult child",
    "stepchild",
    "step child",
    "stepson",
    "stepdaughter",
    "foster child",
    "foster",
    "niece",
    "nephew",
    "ward",
    "godchild",
}
EX_TYPE = re.compile(r"\bex\b|^ex|divorc|former|separated", re.I)

# ---------------------------------------------------------------------------
# Normalization tables

# Alias groups: any two names in the same group can be the same person.
# A short form may sit in several groups (Pat: Patrick and Patricia; Ted:
# Edward and Theodore), so names are compared by group membership rather
# than mapped to a single canonical form, which would make Patricia match
# Patrick through "Pat".
NAME_GROUPS = [
    "william bill billy will willie liam",
    "robert bob bobby rob robbie bert",
    "richard dick rick ricky rich richie",
    "james jim jimmy",
    "joseph joe joey jo",
    "thomas tom tommy",
    "michael mike mikey mickey",
    "david dave davey",
    "daniel dan danny",
    "charles chuck charlie chick chas",
    "edward ed eddie ted teddy ned",
    "theodore ted teddy theo",
    "anthony antonio tony",
    "antoinette toni",
    "stephen steven steve stevie",
    "christopher chris kit",
    "christine christina chris chrissy tina",
    "matthew matt",
    "patrick pat",
    "patricia pat patty patsy tricia trish",
    "katherine catherine kathryn kathleen kathy cathy kate katie kat",
    "elizabeth beth liz lizzie betty betsy eliza libby",
    "margaret peggy maggie meg marge margie",
    "susan sue susie suzanne",
    "deborah debra debbie deb",
    "jennifer jenny jen jenn",
    "barbara barb barbie",
    "teresa theresa therese terry terri tess tessa",
    "terrence terence terry",
    "cynthia cindy",
    "sandra sandy",
    "pamela pam",
    "rebecca becky becca",
    "gerald jerry gerry",
    "lawrence laurence larry",
    "kenneth ken kenny",
    "ronald ron ronnie",
    "donald don donnie",
    "gregory greg",
    "jeffrey geoffrey jeff",
    "timothy tim timmy",
    "andrew andy drew",
    "benjamin ben benny",
    "samuel sam sammy",
    "samantha sam sammy",
    "alexander alex",
    "alexandra alex",
    "nicholas nick nicky",
    "frederick fred freddie",
    "francis frank fran",
    "frances fran",
    "eugene gene",
    "henry hank",
    "john jack johnny jon",
    "jonathan jon",
    "raymond ray",
    "walter walt",
    "phillip philip phil",
    "peter pete",
    "vincent vince",
    "joshua josh",
    "jacob jake",
    "zachary zach zack",
    "nathan nathaniel nate",
    "sarah sara sally",
    "abigail abby",
    "amanda mandy",
    "victoria vicky vickie tori",
    "virginia ginny ginger",
    "dorothy dottie dot",
    "judith judy",
    "angela angie",
    "joan joanie jo",
    "joanne joann jo",
    "josephine jo",
    "johanna jo",
    "mary marie maria ma molly mamie",
    "ann anne anna ana annie",
    "carol carole",
    "leonard leo len lenny",
    "albert al bert",
    "alan allen allan al",
    "arthur art",
    "harold hal harry",
    "douglas doug",
    "kimberly kim",
    "rosemary rose rosie",
    "jacqueline jacquelyn jackie",
    "jessica jess",
]
NAME_GROUP_IDS = collections.defaultdict(set)
for _gid, _names in enumerate(NAME_GROUPS):
    for _n in _names.split():
        NAME_GROUP_IDS[_n].add(_gid)


def name_keys(n):
    """Blocking keys for a given name: its first word's alias groups, or that word."""
    n = n.split()[0]
    return {f"g{gid}" for gid in NAME_GROUP_IDS[n]} if n in NAME_GROUP_IDS else {n}


# Similar-looking names that belong to different sexes; never a match.
CROSS_GENDER = {
    frozenset(p.split())
    for p in [
        "patricia patrick",
        "michael michelle",
        "christine christopher",
        "christina christopher",
        "frances francis",
        "stephanie stephen",
        "gerald geraldine",
        "paul paula",
        "paul pauline",
        "joseph josephine",
        "daniel danielle",
        "eric erica",
        "carl carla",
        "robert roberta",
        "louis louise",
        "henry henrietta",
        "george georgia",
        "charles charlotte",
        "andrew andrea",
        "albert alberta",
        "antonio antonia",
        "mario maria",
        "julius julia",
        "julian juliana",
        "dean deanna",
        "fred freda",
        "bernard bernadette",
        "glen glenda",
        "gene genevieve",
        "jean john",
        "jan janet",
        "ronald ronda",
        "brian briana",
        "kevin kevina",
        "nicholas nicole",
        "victor victoria",
        "justin justine",
        "gabriel gabriella",
        "regina reginald",
        "jose josefa",
        "tony tonya",
        "emil emily",
        "jesse jessica",
    ]
}

STREET_WORDS = {
    "road": "rd",
    "street": "st",
    "drive": "dr",
    "avenue": "ave",
    "av": "ave",
    "lane": "ln",
    "court": "ct",
    "circle": "cir",
    "boulevard": "blvd",
    "place": "pl",
    "parkway": "pkwy",
    "highway": "hwy",
    "terrace": "ter",
    "trail": "trl",
    "way": "wy",
    "square": "sq",
    "point": "pt",
    "pike": "pk",
    "north": "n",
    "south": "s",
    "east": "e",
    "west": "w",
    "northeast": "ne",
    "northwest": "nw",
    "southeast": "se",
    "southwest": "sw",
    "saint": "st",
    "mount": "mt",
    "fort": "ft",
    "apartment": "unit",
    "apt": "unit",
    "suite": "unit",
    "ste": "unit",
    "lot": "unit",
    "number": "unit",
    "no": "unit",
    "post": "po",
    "office": "",
}
# Street types dropped when they are the last word of the street part
# ("Parkway" == "Parkway Dr"); never dropped from the middle ("St Matthews").
STREET_TYPES = {
    "rd",
    "st",
    "dr",
    "ave",
    "ln",
    "ct",
    "cir",
    "blvd",
    "pl",
    "pkwy",
    "hwy",
    "ter",
    "trl",
    "wy",
    "sq",
    "pk",
}
SUFFIXES = {"jr", "sr", "ii", "iii", "iv"}
TITLES = {"mr", "mrs", "ms", "miss", "dr", "rev", "fr", "deacon", "dcn"}
# Surname tokens too generic to show two surnames are related.
SURNAME_STOP = {
    "estate",
    "estates",
    "of",
    "the",
    "funeral",
    "home",
    "family",
    "trust",
    "church",
    "and",
    "van",
    "von",
    "vander",
    "del",
    "de",
    "der",
    "den",
    "la",
    "le",
    "mc",
    "mac",
    "st",
    "saint",
    "di",
    "da",
    "du",
    "dos",
    "das",
    "y",
}
PLACEHOLDER_LAST = re.compile(r"loose|unknown family|anonymous|certificate|^no name$|^unknown$|^cash$", re.I)
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[a-z]{2,}$")
PHONE_RE = re.compile(r"\(?(\d{3})\)?[\s.\-]*(\d{3})[\s.\-]*(\d{4})")


def text(v):
    """Return a stripped string for a cell value ('' for empty/placeholder)."""
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("(blank)", "none", "null") else s


def norm_last(v):
    """Lowercase surname without punctuation or generational suffix.

    Apostrophes are deleted (O'Brien == OBrien); other punctuation splits.
    """
    s = re.sub(r"['’`]", "", text(v).lower())
    s = re.sub(r"[^a-z\- ]", " ", s)
    toks = [t for t in s.replace("-", " - ").split() if t not in SUFFIXES]
    return " ".join(toks).replace(" - ", "-").strip("- ")


def first_names(v):
    """Return the set of given names in a first-names field (lowercase).

    "Chick & Kathy" -> {"chick", "kathy"}; nickname equivalence is decided
    later by name_match.  One name per "&"/"and"/","/"/" part, made of its
    first two words that are not titles or bare initials, so
    "Mary Helen" -> "mary helen" and "J. Robert" -> "robert".
    """
    s = text(v).lower()
    if not s or s == "no name":
        return frozenset()
    raw = set()
    for part in re.split(r"&|\band\b|,|/|\+", s):
        toks = [t for t in re.sub(r"[^a-z ]", " ", part).split() if t not in TITLES]
        toks = [t for t in toks if len(t) > 1] or toks
        if toks:
            raw.add(" ".join(toks[:2]))
    return frozenset(raw)


def suffixes(*vals):
    """Generational suffixes (jr/sr/ii...) mentioned in any of the name fields."""
    toks = set()
    for v in vals:
        toks |= set(re.sub(r"[^a-z ]", " ", text(v).lower()).split())
    return frozenset(toks & SUFFIXES)


def norm_email(v):
    """Lowercased email, or '' when it is not a plausible address."""
    s = text(v).lower()
    return s if EMAIL_RE.match(s) else ""


def norm_phone(v):
    """10-digit phone number, or '' (placeholders and 7-digit numbers dropped).

    The first 3-3-4 digit group wins, so extensions ("x12") and a leading
    country code are ignored.
    """
    m = PHONE_RE.search(re.sub(r"^\s*\+?1[\s.\-]*(?=\(?\d{3})", "", text(v)))
    if not m:
        return ""
    d = "".join(m.groups())
    if len(set(d)) == 1 or d[3:6] == "555" or d[0] in "01":
        return ""
    return d


def norm_zip(v):
    """5-digit ZIP (restores leading zeros lost to numeric cells)."""
    d = text(v).split("-")[0].strip()
    if d.isdigit() and len(d) <= 5:
        return d.zfill(5)
    if d.isdigit() and len(d) in (8, 9):
        return d.zfill(9)[:5]
    return d[:5] if d[:5].isdigit() else ""


def norm_addr(a1, a2):
    """Normalized street address, or '' when there is no street number/box.

    Standardizes directionals, Saint/Mount, P.O. Box and unit designators,
    and drops a trailing street type, so "123 North Main Street, Apt. 4"
    == "123 N Main #4" and "126 N. Birchwood Avenue" == "126 N Birchwood".
    """
    s = (text(a1) + " " + text(a2)).lower().replace("#", " unit ")
    toks = [STREET_WORDS.get(t, t) for t in re.sub(r"[^a-z0-9 ]", " ", s).split()]
    toks = [t for t in toks if t]
    # "p o box" -> "po box"; collapse "unit unit" from e.g. "Apt #4".
    out = []
    for t in toks:
        if out and out[-1] == "p" and t == "o":
            out[-1] = "po"
        elif not (t == "unit" and out and out[-1] == "unit"):
            out.append(t)
    unit = out.index("unit") if "unit" in out else len(out)
    street, rest = out[:unit], out[unit:]
    if len(street) >= 3 and street[-1] in STREET_TYPES:
        street = street[:-1]
    key = " ".join(street + rest)
    return key if re.search(r"\d", key) else ""


def as_date(v):
    """Coerce a cell to datetime or None."""
    return v if isinstance(v, datetime.datetime) else None


def to_int(v):
    """Integer DUID from a cell, or None for malformed values."""
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Loading


def cached(fn):
    """Cache fn(path)'s result in cache_dir; reloading the XLSX takes minutes.

    The cache name includes this script's name (so sibling scripts never
    share a cache), it is reused only if newer than both the XLSX and this
    script, and it is written atomically so an interrupted run cannot leave
    a truncated "fresh" cache.
    """

    def wrapper(path, cache_dir):
        cache = Path(cache_dir) / f"{Path(path).stem}.{Path(__file__).stem}.cache.pickle"
        newest = max(Path(path).stat().st_mtime, Path(__file__).stat().st_mtime)
        if cache.exists() and cache.stat().st_mtime > newest:
            with open(cache, "rb") as fh:
                return pickle.load(fh)
        result = fn(path)
        tmp = cache.with_suffix(".tmp")
        with open(tmp, "wb") as fh:
            pickle.dump(result, fh)
        os.replace(tmp, cache)
        return result

    return wrapper


def read_sheet(path, cols):
    """Yield ({key: value}, raw row) for the first sheet, using cols {key: header}."""
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = wb.worksheets[0].iter_rows(values_only=True)
    hdr = list(next(rows))
    idx = {k: hdr.index(v) for k, v in cols.items()}
    for r in rows:
        yield {k: r[i] for k, i in idx.items()}, r


def assign_roles(ms):
    """Set each member's "role" (partner/child/ex/other) within one Family.

    Explicit types map directly.  Ambiguous (Unknown/Adult/...) members are
    children if under 18 today or if another member is GENERATION_YEARS+
    older; "other" if an explicit head/spouse is GENERATION_YEARS+ younger
    (a parent living with the household) or, lacking a birthdate, if the
    Family already has two explicit heads/spouses; otherwise partners.  A
    partner assigned by default with no birthdate is flagged "guessed".
    """
    births = [m["birth"] for m in ms if m["birth"]]
    explicit = [m["birth"] for m in ms if m["birth"] and m["type"].lower() in PARTNER_TYPES]
    n_explicit = sum(m["type"].lower() in PARTNER_TYPES for m in ms)
    for m in ms:
        m["guessed"] = False
        t = m["type"].lower()
        parent_word = re.search(r"mother|father|parent", t)
        if EX_TYPE.search(t):
            m["role"] = "ex"
        elif t in PARTNER_TYPES or t.startswith("dec"):
            m["role"] = "partner"
        elif t in CHILD_TYPES or (not parent_word and t.startswith(CHILD_PREFIXES)):
            m["role"] = "child"
        elif t in AMBIGUOUS_TYPES:
            b = m["birth"]
            minor = b and (TODAY - b).days < 18 * 365.25
            older = b and any(o.year <= b.year - GENERATION_YEARS for o in births)
            parent = b and any(o.year >= b.year + GENERATION_YEARS for o in explicit)
            # With no birthdate, and an explicit couple already heading the
            # Family, an Unknown/Adult member is someone else (often a child).
            couple = not b and n_explicit >= 2
            m["role"] = "child" if minor or older else "other" if parent or couple else "partner"
            m["guessed"] = m["role"] == "partner" and not b
        else:
            m["role"] = "other"


@cached
def load_members(path):
    """Load the Member export as {family DUID: [member dict, ...]}.

    Each member carries normalized given names (formal + nickname), aliases
    (also the middle name, for people who go by it), surnames (last +
    maiden), birthdate and whether it is precise (Jan 1 is usually a
    year-only placeholder), role, suffix, deceased flag, contact info, and
    Date Modified unless it was a bulk update.
    """
    raw = [m for m, _ in read_sheet(path, MEMBER_COLS)]
    per_minute = collections.Counter(
        as_date(m["modified"]).strftime("%Y%m%d%H%M") for m in raw if as_date(m["modified"])
    )
    members = collections.defaultdict(list)
    for m in raw:
        # Contribution-record placeholders put the surname in the first-name
        # field; they are not people and would look like extra partners.
        # A surname-as-first-name entry with a birthdate, a distinct
        # nickname, or a specific type (Daughter, Husband...) is a real person.
        if (
            norm_last(m["first"])
            and norm_last(m["first"]) == norm_last(m["last"])
            and not as_date(m["birth"])
            and norm_last(m["nick"]) in ("", norm_last(m["last"]))
            and text(m["type"]).lower() in ("", "head", "unknown")
        ):
            continue
        firsts = first_names(m["first"]) | first_names(m["nick"])
        birth = as_date(m["birth"])
        mod = as_date(m["modified"])
        if mod and per_minute[mod.strftime("%Y%m%d%H%M")] >= BULK_MODIFIED_PER_MINUTE:
            mod = None
        mtype = text(m["type"])
        deceased = (
            str(m["deceased"]) == "1"
            or bool(as_date(m["death"]))
            or text(m["status"]).lower() == "deceased"
            or "dec" in mtype.lower()
        )
        members[to_int(m["fam"])].append(
            {
                "duid": m["duid"],
                "type": mtype,
                "firsts": frozenset(firsts),
                "aliases": frozenset(firsts | first_names(m["middle"])),
                "last": norm_last(m["last"]),
                "lasts": frozenset(x for x in (norm_last(m["last"]), norm_last(m["maiden"])) if x),
                "birth": birth.date() if birth else None,
                "precise": bool(birth) and not (birth.month == 1 and birth.day == 1),
                "suffix": re.sub(r"[^a-z]", "", text(m["suffix"]).lower()),
                "deceased": deceased,
                "email": norm_email(m["email"]),
                "cell": norm_phone(m["cell"]),
                "modified": mod,
                "label": f"{text(m['first'])} {text(m['last'])} ({mtype or '?'}"
                + (f", b. {birth:%Y-%m-%d}" if birth else "")
                + (", deceased" if deceased else "")
                + ")",
            }
        )
    for ms in members.values():
        assign_roles(ms)
    return dict(members)


@cached
def load_enhanced(path):
    """Load extra Family fields from the enhanced export: {DUID: dict}.

    That export has rows broken by multi-line Notes (values shifted into the
    wrong columns), so only rows whose DUID parses and whose "is Registered"
    is Yes/No are used.  Repeated DUIDs keep their first sane row.
    """
    out = {}
    for e, _ in read_sheet(path, ENHANCED_COLS):
        duid = to_int(e["duid"])
        if duid is None or e["registered"] not in ("Yes", "No") or duid in out:
            continue
        desc = text(e["emergency_desc"])
        phone = norm_phone(e["emergency"])
        out[duid] = {
            "last_updated": as_date(e["last_updated"]),
            "emergency": "" if NOT_OWN_PHONE.search(desc) else phone,
        }
    return out


@cached
def load(path):
    """Load and normalize the Family export."""
    fams = []
    for f, r in read_sheet(path, COLS):
        if f["duid"] is None:
            continue
        f.update(
            n_last=norm_last(f["last"]),
            n_first=first_names(f["first"]),
            n_suffix=suffixes(f["first"], f["last"], f["mailing"]),
            n_email=norm_email(f["email"]),
            n_phones=frozenset(p for p in (norm_phone(f["home_phone"]), norm_phone(f["primary_phone"])) if p),
            n_zip=norm_zip(f["zip"]),
            n_city=text(f["city"]).lower(),
            filled=sum(1 for v in r if text(v)),
        )
        f["n_addr"] = norm_addr(f["addr1"], f["addr2"])
        if f["n_addr"]:
            f["n_addr"] += " | " + (f["n_zip"] or f["n_city"])
        # Placeholders, plus a few rows whose columns are shifted (a phone
        # number sitting in the ZIP column), are excluded from matching.
        f["placeholder"] = bool(
            not f["n_last"]
            or PLACEHOLDER_LAST.search(text(f["last"]))
            or "loose" in text(f["mailing"]).lower()
            or len(re.sub(r"\D", "", text(f["zip"]))) > 9
        )
        fams.append(f)
    return fams


def bulk_days(dated):
    """Bulk keys from (parish, datetime) pairs.

    Returns a set holding each calendar day with BULK_DAY_MIN+ dates
    Archdiocese-wide, and each (parish, day) with BULK_PARISH_DAY_MIN+.
    """
    dated = [(org, d.date()) for org, d in dated if d]
    days = collections.Counter(day for _, day in dated)
    parish_days = collections.Counter(dated)
    return {day for day, n in days.items() if n >= BULK_DAY_MIN} | {
        key for key, n in parish_days.items() if n >= BULK_PARISH_DAY_MIN
    }


def is_bulk(f, d, bulk):
    """True if Family f's datetime d falls on a bulk day (global or its parish)."""
    return d.date() in bulk or (f["org_id"], d.date()) in bulk


def enrich(fams, members, enhanced):
    """Attach members and enhanced fields; derive contacts and activity.

    Mutates the Family dicts after they were cached, so the cache stays raw.
    Returns the bulk-day sets used, for the report.
    """
    extra = {f["duid"]: enhanced.get(f["duid"], {}) for f in fams}
    unreg_bulk = bulk_days((f["org_id"], as_date(f["unregistered"])) for f in fams)
    upd_bulk = bulk_days((f["org_id"], extra[f["duid"]].get("last_updated")) for f in fams)
    for f in fams:
        ms = members.get(f["duid"], [])
        e = extra[f["duid"]]
        f["members"] = ms
        f["n_emails"] = frozenset({f["n_email"]} | {m["email"] for m in ms}) - {""}
        f["n_phones"] = frozenset(f["n_phones"] | {m["cell"] for m in ms} | {e.get("emergency", "")}) - {""}
        # No member data: unknown, so not ranked like an all-deceased Family.
        f["living"] = any(not m["deceased"] for m in ms) if ms else True
        f["catchall"] = len({m["last"] for m in ms if m["last"]}) >= CATCHALL_SURNAMES
        f["last_updated"] = e.get("last_updated")
        # Most recent activity: membership date, non-bulk unregistration,
        # non-bulk Family "last updated", non-bulk Member "Date Modified".
        dates = [as_date(f["membership"])]
        unreg = as_date(f["unregistered"])
        if unreg and not is_bulk(f, unreg, unreg_bulk):
            dates.append(unreg)
        if f["last_updated"] and not is_bulk(f, f["last_updated"], upd_bulk):
            dates.append(f["last_updated"])
        dates += [m["modified"] for m in ms if m["modified"]]
        dates = [d for d in dates if d]
        f["activity"] = max(dates) if dates else datetime.datetime.min
    return unreg_bulk, upd_bulk


# ---------------------------------------------------------------------------
# Matching


def shared_values(fams, key):
    """Contact values used by more than MAX_SURNAMES_PER_CONTACT distinct surnames."""
    surnames = collections.defaultdict(set)
    for f in fams:
        vals = f[key] if isinstance(f[key], frozenset) else [f[key]]
        for v in vals:
            if v:
                surnames[v].add(f["n_last"])
    return {v for v, s in surnames.items() if len(s) > MAX_SURNAMES_PER_CONTACT}


def candidate_pairs(fams):
    """Index pairs that share an email, phone, address, surname+ZIP, or person.

    Blocking avoids comparing all ~159k^2/2 pairs.  Returns (pairs, number
    of oversized blocks skipped).
    """
    blocks = collections.defaultdict(list)
    for i, f in enumerate(fams):
        if f["placeholder"]:
            continue
        for e in f["n_emails"]:
            blocks["e:" + e].append(i)
        for p in f["n_phones"]:
            blocks["p:" + p].append(i)
        if f["n_addr"]:
            blocks["a:" + f["n_addr"]].append(i)
        if f["n_zip"]:
            blocks["lz:" + f["n_last"] + "|" + f["n_zip"]].append(i)
        # Same person (given name + exact birthdate + surname or maiden
        # name) anywhere in the Archdiocese: catches moves where every
        # Family-level field changed.
        for m in [] if f["catchall"] else f["members"]:
            if m["precise"]:
                for key in {k for first in m["firsts"] for k in name_keys(first)}:
                    for last in m["lasts"]:
                        blocks[f"m:{key}|{m['birth']}|{last}"].append(i)
    pairs, skipped = set(), 0
    for idxs in blocks.values():
        idxs = sorted(set(idxs))
        if len(idxs) > MAX_BLOCK:
            skipped += 1
            continue
        for a in range(len(idxs)):
            for b in range(a + 1, len(idxs)):
                pairs.add((idxs[a], idxs[b]))
    return pairs, skipped


def last_relation(a, b):
    """'same', 'similar' (typo or compound-surname overlap), or None.

    Compound overlap needs a shared token of 4+ letters that is not a
    particle or generic word ("van", "del", "estate", "funeral", ...).
    """
    if not (a and b):
        return None
    if a == b:
        return "same"
    ta, tb = set(re.split(r"[- ]", a)), set(re.split(r"[- ]", b))
    if any(len(t) >= 4 and t not in SURNAME_STOP for t in ta & tb):
        return "similar"
    if min(len(a), len(b)) >= 4 and JaroWinkler.similarity(a, b) >= SIMILAR_LAST:
        return "similar"
    return None


def name_match(a, b):
    """True if two given names (one or two words) can be the same person's.

    Two-word names: the first words must match, and if both have a second
    word, those must match too ("Mary Helen" != "Mary Margie", but
    "Mary Helen" == "Mary").
    """
    a, b = a.split(), b.split()
    if "".join(a) == "".join(b):  # "La Shauna" == "LaShauna"
        return True
    if not word_match(a[0], b[0]):
        return False
    if len(a) < 2 or len(b) < 2:
        return True
    x, y = sorted((a[1], b[1]), key=len)
    return word_match(x, y) or (len(x) >= 3 and y.startswith(x))  # "Mary Lou" == "Mary Louise"


def gendered_pair(a, b):
    """True for masculine/feminine forms of one name (Juan/Juana, Alejandro/Alejandra)."""
    s, lng = sorted((a, b), key=len)
    if frozenset((a, b)) in CROSS_GENDER:
        return True
    if len(s) >= 3 and lng == s + "a":
        return True
    return len(a) == len(b) >= 4 and a[:-1] == b[:-1] and {a[-1], b[-1]} == {"o", "a"}


def word_match(a, b):
    """True if two single given-name words can be the same person's name.

    Matches: the same alias group (Bill/William, Pat/Patricia, Ted/Theodore);
    a 4+ letter name fully prefixing a longer one by up to 3 letters
    (Joan/Joanie/Joanna); close Jaro-Winkler similarity on longer names.
    Masculine/feminine forms (Patricia/Patrick, Juan/Juana) never match
    unless an alias group says so (Ann/Anna).
    """
    if a == b:
        return True
    if NAME_GROUP_IDS.get(a, set()) & NAME_GROUP_IDS.get(b, set()):
        return True
    if gendered_pair(a, b):
        return False
    s, lng = sorted((a, b), key=len)
    if len(s) >= 4 and lng.startswith(s) and len(lng) - len(s) <= 3:
        return True
    return len(s) >= 5 and JaroWinkler.similarity(a, b) >= 0.93


def member_status(m, n, aliases=False):
    """Compare two members: 'strong', 'year', 'name', 'conflict', or None.

    None: different names.  'strong': same name and same precise birthdate.
    'year': same name, birthdates agree as far as they go (year-only
    placeholder, or a one-field typo).  'name': same name, a birthdate is
    missing.  'conflict': same name but clearly different people
    (different birthdates or Jr. vs Sr.).  With aliases=True, middle names
    count as given names (used only to decide whether a partner is present
    in the other Family, never as identity evidence).
    """
    key = "aliases" if aliases else "firsts"
    if not any(name_match(a, b) for a in m[key] for b in n[key]):
        return None
    if not (m["lasts"] & n["lasts"] or last_relation(m["last"], n["last"])):
        return None
    if m["suffix"] and n["suffix"] and m["suffix"] != n["suffix"]:
        return "conflict"
    a, b = m["birth"], n["birth"]
    if not (a and b):
        return "name"
    if a == b:
        return "strong" if m["precise"] and n["precise"] else "year"
    if not (m["precise"] and n["precise"]):
        return "year" if a.year == b.year else "conflict"
    # Typos: same year with one of month/day mistyped or month/day swapped;
    # or same month/day with a year that differs in one digit or by two
    # swapped digits.  Anything more (e.g. 1954 vs 1979: father and son
    # sharing a birthday) is two different people.
    if a.year == b.year:
        swapped = a.month == b.day and a.day == b.month
        return "year" if a.month == b.month or a.day == b.day or swapped else "conflict"
    if (a.month, a.day) == (b.month, b.day):
        ya, yb = str(a.year), str(b.year)
        diff = [k for k in range(4) if ya[k] != yb[k]]
        if len(diff) == 1 or (len(diff) == 2 and ya[diff[0]] == yb[diff[1]] and ya[diff[1]] == yb[diff[0]]):
            return "year"
    return "conflict"


PRESENT = ("strong", "year", "name")


def member_evidence(f, g):
    """Summarize member-level evidence between two Families (symmetric).

    Returns a dict:
      usable          both Families have members and neither is a catch-all
      strong_partner  partner<->partner pairs matching by name + exact
                      birthdate (each member used at most once)
      strong_child    same-role non-partner pairs (child<->child, ...)
      role_mismatch   a person matching by name (any birthdate agreement)
                      who is a partner in one Family and a child/other in
                      the other (grown child; parent moved in)
      ex_link         the shared person is an ex-spouse in exactly one
                      Family (the ex's own household)
      soft_partner    a partner pair matches by name with consistent or
                      missing birthdates
      conflict        a partner pair has the same name but clearly different
                      birthdates/suffixes, and neither has a better match
      partners_shared some partner is present in both Families
      living_shared   ... and is living in both
      lenient_shared  some partner of each Family is present in the other
                      by the lenient presence rules
      living_f/_g     living partners in f (g)
      shared_births   known birthdates of living partners in both Families
      *_list_f/_g     the living / extra / absent-deceased partner members
      extra_f/extra_g living partners of f (g) absent from g (f)
      dead_f/dead_g   deceased partners of f (g) absent from g (f)
      any_match       any member pair matches by name
    """
    ev = {
        "usable": False,
        "strong_partner": 0,
        "strong_child": 0,
        "role_mismatch": False,
        "ex_link": False,
        "soft_partner": False,
        "conflict": False,
        "partners_shared": False,
        "extra_f": 0,
        "extra_g": 0,
        "dead_f": 0,
        "dead_g": 0,
        "any_match": False,
        "living_shared": False,
        "living_f": 0,
        "living_g": 0,
        "lenient_shared": False,
        "shared_births": [],
        "living_list_f": [],
        "living_list_g": [],
        "extra_list_f": [],
        "extra_list_g": [],
        "dead_list_f": [],
        "dead_list_g": [],
    }
    F, G = f["members"], g["members"]
    if not F or not G or f["catchall"] or g["catchall"]:
        return ev
    ev["usable"] = True
    st = {}
    for i, m in enumerate(F):
        for j, n in enumerate(G):
            v = member_status(m, n)
            if v:
                st[(i, j)] = v
    ev["any_match"] = any(v in PRESENT for v in st.values())

    # One-to-one exact matches, preferring same-role pairings, ordered
    # symmetrically (by the pair's member DUIDs) so f/g order cannot matter.
    strong = sorted(
        (k for k, v in st.items() if v == "strong"),
        key=lambda k: (F[k[0]]["role"] != G[k[1]]["role"], sorted((str(F[k[0]]["duid"]), str(G[k[1]]["duid"])))),
    )
    used_i, used_j = set(), set()
    for i, j in strong:
        if i in used_i or j in used_j:
            continue
        used_i.add(i)
        used_j.add(j)
        ri, rj = F[i]["role"], G[j]["role"]
        if ri == rj == "partner":
            ev["strong_partner"] += 1
        elif "partner" not in (ri, rj) and "ex" not in (ri, rj):
            ev["strong_child"] += 1

    for (i, j), v in st.items():
        if v not in PRESENT:
            continue
        ri, rj = F[i]["role"], G[j]["role"]
        roles = {ri, rj}
        if "ex" in roles and ri != rj:
            ev["ex_link"] = True
        elif "partner" in roles and roles & {"child", "other"}:
            ev["role_mismatch"] = True
        elif ri == rj == "partner":
            ev["partners_shared"] = True
            if not (F[i]["deceased"] or G[j]["deceased"]):
                ev["living_shared"] = True
            if v != "strong":
                ev["soft_partner"] = True

    # Same-named partners with different birthdates are different people,
    # unless either has a better partner match, or either's exact birthdate
    # belongs to a partner on the other side (spouses' birthdates swapped
    # between the two records).
    f_ok = {i for (i, j), v in st.items() if v in PRESENT and F[i]["role"] == G[j]["role"] == "partner"}
    g_ok = {j for (i, j), v in st.items() if v in PRESENT and F[i]["role"] == G[j]["role"] == "partner"}
    f_births = {m["birth"] for m in F if m["role"] == "partner" and m["precise"]}
    g_births = {n["birth"] for n in G if n["role"] == "partner" and n["precise"]}
    ev["conflict"] = any(
        v == "conflict"
        and F[i]["role"] == G[j]["role"] == "partner"
        and i not in f_ok
        and j not in g_ok
        and not (F[i]["precise"] and F[i]["birth"] in g_births)
        and not (G[j]["precise"] and G[j]["birth"] in f_births)
        for (i, j), v in st.items()
    )

    # Partner presence is deliberately lenient, so a spouse recorded
    # differently in the two records (maiden name, typo, swapped names,
    # goes by middle name) is not mistaken for a different spouse.  A
    # partner m is present in the other Family if it has: the same person
    # exactly (any role); or a partner/ex with the same name (middle names
    # allowed), the same precise birthdate, or -- when the two Families'
    # surnames are related -- the same given name (or a one-letter
    # variant) under any surname; or m's given name appears in the other
    # Family's first-names field.  Only partners/exes count for the lenient
    # rules (a child named after a parent must not stand in for an absent
    # spouse), and anyone whose birthdate conflicts is skipped.
    related = last_relation(f["n_last"], g["n_last"]) is not None

    def variant(a, b):
        """Given names that match, differ by one letter (4+ letters), or where
        one person goes by the other record's second given name."""
        if name_match(a, b):
            return True
        wa, wb = a.split(), b.split()
        if min(len(wa[0]), len(wb[0])) >= 4 and Levenshtein.distance(wa[0], wb[0]) <= 1:
            return True
        return any(word_match(x, y) for x in wa[1:] for y in wb[:1]) or any(
            word_match(x, y) for x in wa[:1] for y in wb[1:]
        )

    def present(m, B, other_first_names):
        for n in B:
            alias_status = member_status(m, n, aliases=True)
            if alias_status == "conflict":
                continue
            if member_status(m, n) == "strong":
                return True
            if n["role"] not in ("partner", "ex"):
                # A partner by default (ambiguous type, no birthdate) may
                # really be the child this name matches.
                if m["guessed"] and alias_status in PRESENT:
                    return True
                continue
            if alias_status in PRESENT:
                return True
            if m["precise"] and n["precise"] and m["birth"] == n["birth"]:
                return True
            births_differ = m["birth"] and n["birth"] and m["birth"] != n["birth"]
            if related and not births_differ and any(variant(a, b) for a in m["aliases"] for b in n["aliases"]):
                return True
        return related and any(name_match(a, b) for a in m["firsts"] for b in other_first_names)

    present_count = {"f": 0, "g": 0}
    for A, B, side, other in ((F, G, "f", g), (G, F, "g", f)):
        for m in A:
            if m["role"] != "partner":
                continue
            if not m["deceased"]:
                ev["living_" + side] += 1
                ev["living_list_" + side].append(m)
            if present(m, B, other["n_first"]):
                present_count[side] += 1
            else:
                kind = "dead_" if m["deceased"] else "extra_"
                ev[kind + side] += 1
                ev[kind + "list_" + side].append(m)
    # Some partner of each Family is present in the other, by the lenient
    # rules (covers spouses recorded differently).
    ev["lenient_shared"] = ev["partners_shared"] or (present_count["f"] and present_count["g"]) > 0
    # Known birthdate of each living partner shared by both Families (from
    # whichever copy has one), for the widowed generation check.
    ev["shared_births"] = [
        F[i]["birth"] or G[j]["birth"]
        for (i, j), v in st.items()
        if v in PRESENT
        and F[i]["role"] == G[j]["role"] == "partner"
        and not (F[i]["deceased"] or G[j]["deceased"])
        and (F[i]["birth"] or G[j]["birth"])
    ]
    return ev


def widowed(ev):
    """Widowed and remarried: one Family has exactly one extra living partner
    (the new spouse), the other has an absent deceased partner, a living
    partner is in both, and the deceased and new spouses are both of the
    shared partner's generation (within GENERATION_YEARS, where known).
    Otherwise e.g. a widowed parent and a grown child's household would
    look like one remarried person."""
    side = "g" if ev["extra_list_g"] else "f"
    other = "f" if side == "g" else "g"
    extras, deads = ev["extra_list_" + side], ev["dead_list_" + other]
    if len(extras) != 1 or not deads or not ev["living_shared"]:
        return False

    def near(a, b):
        return not (a and b) or abs(a.year - b.year) < GENERATION_YEARS

    return all(near(x["birth"], b) for b in ev["shared_births"] for x in deads + extras)


def member_verdict(ev):
    """Return the reason member evidence shows two Families are related or
    different people (never merged or grouped together), else None.

    Shared by score_pair and by grouping's check of pairs that were never
    scored.  "No adults in common" is handled separately in score_pair.
    """
    if not ev["usable"]:
        return None
    if ev["ex_link"] and not ev["partners_shared"]:
        return "related, not duplicate: shared person is an ex-spouse in one Family"
    if ev["role_mismatch"] and not (ev["partners_shared"] or ev["strong_child"]):
        return (
            "related, not duplicate: head/spouse in one Family is a child/other in the other "
            "(grown child, or parent moved in)"
        )
    if ev["extra_f"] and ev["extra_g"] and ev["lenient_shared"]:
        return (
            "related, not duplicate: same person with a different living partner in each Family (divorce/remarriage?)"
        )
    if ev["conflict"] and not (ev["partners_shared"] or ev["strong_child"]):
        return "vetoed: same-named adults with different birthdates or Jr./Sr."
    return None


def suffix_conflict(f, g):
    """Jr. vs Sr. (or II vs III) in the two Families' names: two people."""
    return bool(f["n_suffix"] and g["n_suffix"] and not (f["n_suffix"] & g["n_suffix"]))


def score_pair(f, g, shared, stats):
    """Return (tier, reasons) for a match, None for no match, or
    (None, reason) for a pair that member evidence shows are related or
    different people (kept apart when grouping).

    Tier 1 = very high confidence ... tier 4 = low/review only.  `reasons` is a
    tuple of short phrases that become the human-readable reason code.
    Family-level rules give a base tier; member evidence then confirms,
    upgrades, downgrades, vetoes, or reclassifies the pair as related
    (not duplicate).  `stats` counts vetoes/exclusions for the report.
    """
    # Jr. vs Sr. (or II vs III) on otherwise-identical records: two people.
    if suffix_conflict(f, g):
        return None, "vetoed: Jr./Sr. in Family names"

    last = last_relation(f["n_last"], g["n_last"])
    fa, ga = f["n_first"], g["n_first"]
    first_known = bool(fa and ga)
    first_exact = (
        first_known
        and all(any(name_match(a, b) for b in ga) for a in fa)
        and all(any(name_match(a, b) for a in fa) for b in ga)
    )
    first_overlap = first_known and any(name_match(a, b) for a in fa for b in ga)
    nickname = first_overlap and not (fa & ga)

    email = bool((f["n_emails"] & g["n_emails"]) - shared["e"])
    phone = bool((f["n_phones"] & g["n_phones"]) - shared["p"])
    addr_same = bool(f["n_addr"]) and f["n_addr"] == g["n_addr"]
    addr = addr_same and f["n_addr"] not in shared["a"]
    shared_addr = addr_same and not addr
    contacts = email + phone + addr
    same_zip = bool(f["n_zip"]) and f["n_zip"] == g["n_zip"]
    couple = min(len(fa), len(ga)) >= 2

    # Family-level base tier.
    tier = None
    if last == "same":
        if first_overlap and (contacts >= 2 or (first_exact and contacts >= 1)):
            tier = 1
        elif first_overlap and contacts == 1:
            tier = 2
        elif first_exact and (shared_addr or (couple and same_zip)):
            tier = 3
        elif not first_known and contacts >= 1:
            tier = 3 if contacts == 1 else 2
        elif first_known and not first_overlap and contacts >= 1:
            tier = 4
        elif first_exact and same_zip:
            tier = 4
    elif last == "similar":
        if first_exact and contacts >= 1:
            tier = 2
        elif first_overlap and contacts >= 1:
            tier = 3
    elif first_overlap and (email or phone):
        tier = 4

    # Member-level adjustments, in priority order.  Exclusions are only
    # counted in the report when the Family-level rules alone would have
    # kept the pair (base tier set).
    ev = member_evidence(f, g)

    def excluded(reason):
        # (None, reason) marks a related/vetoed pair: never grouped together.
        if tier is not None:
            stats[reason] += 1
        return None, reason

    # Related/different-people verdicts apply even when there is no match
    # evidence, so the pair is still kept apart if third records link them.
    verdict = member_verdict(ev)
    if verdict:
        return excluded(verdict)
    if tier is None and not (ev["strong_partner"] or ev["strong_child"]):
        return None

    member_reason = None
    if ev["usable"]:
        # Each side has a partner (living or deceased) absent from the
        # other and no partner is shared: different households, e.g. a late
        # couple's record vs. their child's household at the same address.
        if (ev["extra_f"] or ev["dead_f"]) and (ev["extra_g"] or ev["dead_g"]) and not ev["lenient_shared"]:
            singles = ev["living_list_f"] + ev["living_list_g"]
            gap = (
                len(singles) == 2
                and all(m["birth"] for m in singles)
                and (abs(singles[0]["birth"].year - singles[1]["birth"].year) >= GENERATION_YEARS)
            )
            if (
                ev["living_f"] == ev["living_g"] == 1
                and not (ev["dead_f"] or ev["dead_g"] or ev["any_match"])
                and not gap
            ):
                # Two single-adult records sharing no one (e.g. each spouse's
                # record from before the marriage): not one Family, but
                # not kept apart either -- a couple's record may absorb both.
                if tier is not None:
                    stats["not a match: two different single adults (both may still merge into a couple's record)"] += 1
                return None
            return excluded("related, not duplicate: no adults in common (different households)")
        n_strong = ev["strong_partner"] + ev["strong_child"]
        if ev["strong_partner"] and n_strong >= 2:
            tier, member_reason = 1, "2+ members match by name and birthdate"
        elif ev["strong_partner"]:
            tier = 1 if (contacts or first_overlap) else (2 if last else 3)
            member_reason = "an adult matches by name and birthdate"
        elif ev["strong_child"]:
            # Children match but no adult is confirmed: keep strong
            # Family-level evidence, otherwise Medium.
            tier = tier if tier and tier <= 3 else 3
            member_reason = "children match by name and birthdate"
        elif ev["soft_partner"]:
            member_reason = "adults match by name (birthdates missing or partial)"
            if tier == 4 and last:
                tier = 3
        elif not ev["any_match"]:
            member_reason = "no members in common"
        extra_one, dead_other = (ev["extra_g"], ev["dead_f"]) if ev["extra_g"] else (ev["extra_f"], ev["dead_g"])
        if extra_one and dead_other and widowed(ev):
            # A living partner is in both Families and the other Family's
            # partner died: widowed and remarried, not divorce, so no
            # ambiguity and no downgrade.
            member_reason = (member_reason + "; " if member_reason else "") + (
                "spouse deceased in one Family, new spouse in the other (widowed and remarried)"
            )
        elif extra_one:
            # One Family has an adult the other lacks: a pre-marriage single
            # record (duplicate) or a post-divorce one (not).  Ambiguous.
            tier = min(4, tier + 1)
            member_reason = (member_reason + "; " if member_reason else "") + (
                "an adult in one Family is not in the other (new spouse, or ex?)"
            )
    if tier is None:
        return None

    reasons = [
        {"same": "same last name", "similar": "similar last name", None: "different last name"}[last],
    ]
    if member_reason:
        reasons.append(member_reason)
    if not first_known:
        reasons.append("first names missing on one record")
    elif first_exact:
        reasons.append("same first names" + (" (via nickname or spelling variant)" if nickname else ""))
    elif first_overlap:
        reasons.append("first names overlap" + (" (via nickname or spelling variant)" if nickname else ""))
    else:
        reasons.append("different first names")
    if email:
        reasons.append("same email")
    if phone:
        reasons.append("same phone")
    if addr:
        reasons.append("same address")
    if shared_addr:
        reasons.append("same multi-family address")
    if not contacts and not shared_addr:
        reasons.append("same ZIP" if same_zip else "no shared contact info")
    return tier, tuple(reasons)


# ---------------------------------------------------------------------------
# Grouping and survivor choice


class UnionFind:
    """Minimal disjoint-set for clustering matched pairs."""

    def __init__(self):
        self.parent = {}

    def find(self, x):
        self.parent.setdefault(x, x)
        while self.parent[x] != x:
            self.parent[x] = self.parent[self.parent[x]]
            x = self.parent[x]
        return x

    def union(self, a, b):
        self.parent[self.find(a)] = self.find(b)


def survivor_key(f):
    """Sort key: higher = better survivor.

    Order agreed with the Archdiocese: Registered, Active group (so a family
    that moved is merged into the parish where it is Active), then has a
    living member, most recent activity (see enrich), most fields filled
    in, then lowest DUID.
    """
    return (
        text(f["registered"]) == "Yes",
        text(f["group"]) == "Active",
        f["living"],
        f["activity"],
        f["filled"],
        -f["duid"],
    )


def edge_strength(reasons):
    """Rank evidence within one tier (lower = stronger) for grouping order."""
    text_ = "; ".join(reasons)
    rank = 0 if "2+ members match" in text_ else 1 if "an adult matches" in text_ else 2
    return rank + (3 if "is not in the other" in text_ else 0)


def build_groups(fams, edges, cannot, apart):
    """Cluster edges into duplicate groups and assign survivor + member confidence.

    Edges are merged strongest first: by tier, then evidence strength
    (exact member matches, no extra-adult downgrade), then the better
    survivor among the two records.  Tier 1-3 edges cluster transitively,
    except that two groups are never joined if that would put a `cannot`
    pair (related Families, e.g. divorce/remarriage, or vetoed different
    people) in one group -- two records each linked to a third must not be
    merged into each other when the members say they are different
    households.  Any other pair that is not itself a match is checked on
    demand with `apart(i, j)`.  Tier 4 (review-only)
    edges only form groups among records not already in a stronger group,
    so weak evidence never pulls unrelated Families into a strong cluster.
    Each member's confidence is the better of its direct link to the
    survivor and its strongest multi-hop path (rated by its weakest link),
    one level weaker.  Returns (groups, weak edges not used, tier 1-3
    edges blocked by a cannot-link).
    """
    uf = UnionFind()
    comp = {}
    no_link = collections.defaultdict(set)
    for i, j in cannot:
        no_link[i].add(j)
        no_link[j].add(i)

    def try_union(i, j):
        """Union i's and j's groups unless that would join a cannot pair."""
        ri, rj = uf.find(i), uf.find(j)
        if ri == rj:
            return True
        ci, cj = comp.get(ri, {i}), comp.get(rj, {j})
        small, big = (ci, cj) if len(ci) <= len(cj) else (cj, ci)
        if any(no_link[x] & big for x in small):
            return False
        if any(apart(x, y) for x in small for y in big):
            return False
        uf.union(i, j)
        comp.pop(ri, None)
        comp.pop(rj, None)
        comp[uf.find(i)] = ci | cj
        return True

    # Two stable sorts: better survivor first, then tier and evidence strength.
    order = sorted(edges.items(), key=lambda kv: (fams[kv[0][0]]["duid"], fams[kv[0][1]]["duid"]))
    order = sorted(
        order,
        key=lambda kv: max(survivor_key(fams[kv[0][0]]), survivor_key(fams[kv[0][1]])),
        reverse=True,
    )
    order.sort(key=lambda kv: (kv[1][0], edge_strength(kv[1][1])))
    blocked = 0
    for (i, j), (tier, _) in order:
        if tier <= 3 and not try_union(i, j):
            blocked += 1
    # try_union's find() also registers records whose union was blocked, so
    # groups come from `comp` (only actually-joined records), not uf.parent.
    strong = set().union(*comp.values()) if comp else set()
    for (i, j), (tier, _) in order:
        if tier == 4 and i not in strong and j not in strong:
            try_union(i, j)  # a blocked weak link is counted in crossing_weak
    comps = {root: sorted(members) for root, members in comp.items()}
    crossing_weak = sum(1 for (i, j), (t, _) in edges.items() if t == 4 and uf.find(i) != uf.find(j))

    adj = collections.defaultdict(dict)
    for (i, j), ev in edges.items():
        adj[i][j] = adj[j][i] = ev

    groups = []
    for members in comps.values():
        mset = set(members)
        surv = max(members, key=lambda m: survivor_key(fams[m]))
        # Widest-path search from the survivor: a path's strength is its
        # weakest link (highest tier number); prefer stronger, then shorter.
        best = {surv: (0, 0)}
        prev = {}
        heap = [(0, 0, surv)]
        while heap:
            b, h, u = heapq.heappop(heap)
            if (b, h) > best[u]:
                continue
            for n, (t, _) in adj[u].items():
                if n not in mset:
                    continue
                cand = (max(b, t), h + 1)
                if n not in best or cand < best[n]:
                    best[n], prev[n] = cand, u
                    heapq.heappush(heap, (cand[0], cand[1], n))
        rows = []
        for m in members:
            if m == surv:
                continue
            bottleneck, hops = best[m]
            _, reasons = adj[m][prev[m]]
            if hops == 1:
                rows.append((m, bottleneck, reasons))
                continue
            tier = min(4, bottleneck + 1)
            reasons = reasons + ("linked through another record in the group",)
            # In a big group, an indirect link is the classic chaining
            # failure, so it is review-only.  Direct links stay trusted
            # (some Families have 20+ parish records).
            if len(members) > MAX_GROUP:
                tier, reasons = 4, reasons + ("large group, possible chaining",)
            direct = adj[m].get(surv)
            if direct and direct[0] <= tier:
                tier, reasons = direct
            rows.append((m, tier, reasons))
        groups.append({"survivor": surv, "members": rows, "tier": max(t for _, t, _ in rows)})
    return groups, crossing_weak, blocked


# ---------------------------------------------------------------------------
# Classification for the summary report


def match_type(s, f):
    """Coarse story for a (survivor, merged) pair, used in the summary report."""
    same_org = s["org_id"] == f["org_id"]
    s_act = text(s["group"]) == "Active" or text(s["registered"]) == "Yes"
    f_act = text(f["group"]) == "Active" or text(f["registered"]) == "Yes"
    if not same_org and s_act and not f_act:
        return "Moved between parishes (active in one, inactive in the other)"
    if not same_org and s_act and f_act:
        return "Active/registered in two parishes at once"
    if not same_org:
        return "Inactive in two parishes"
    if s_act and f_act:
        return "Two active records in the same parish"
    return "Duplicate record in the same parish (at least one inactive)"


# ---------------------------------------------------------------------------
# Output


def fmt_date(v):
    """Short date string for the report ('' for missing/sentinel dates)."""
    d = as_date(v)
    return d.strftime("%Y-%m-%d") if d and d > datetime.datetime.min else ""


def write_sheet(ws, header, rows, widths):
    """Write a header + rows with bold header, freeze pane, filter and widths."""
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    """Run the full pipeline and write the report files."""
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--families", required=True, type=Path, help="Family export (.xlsx)")
    ap.add_argument("--members", required=True, type=Path, help="Member export (.xlsx)")
    ap.add_argument(
        "--enhanced", type=Path, help="optional enhanced Family export (.xlsx): Date Last Updated, Emergency Phone"
    )
    # Required, with no default, so reports (which contain personal data) are
    # never written into the git checkout by accident.
    ap.add_argument(
        "--out-dir", required=True, type=Path, help="report and cache directory (must be outside the git repo)"
    )
    args = ap.parse_args()
    for path in (args.families, args.members, args.enhanced):
        if path and not path.is_file():
            ap.error(f"no such file: {path}")
    out_dir = args.out_dir
    repo = Path(__file__).resolve().parent
    if out_dir.resolve() == repo or repo in out_dir.resolve().parents:
        ap.error(
            f"--out-dir {out_dir} is inside the git repo; reports contain personal data, so use a directory outside it"
        )
    out_dir.mkdir(parents=True, exist_ok=True)

    fams = load(args.families, out_dir)
    members = load_members(args.members, out_dir)
    enhanced = load_enhanced(args.enhanced, out_dir) if args.enhanced else {}
    unreg_bulk, upd_bulk = enrich(fams, members, enhanced)
    n_members = sum(len(v) for v in members.values())
    no_members = sum(1 for f in fams if not f["members"])
    with_enh = sum(1 for f in fams if f["duid"] in enhanced)
    print(
        f"loaded {len(fams)} families ({sum(f['placeholder'] for f in fams)} placeholders excluded), "
        f"{n_members} members ({no_members} families without members), "
        f"enhanced fields for {with_enh} families"
    )

    shared = {
        "e": shared_values(fams, "n_emails"),
        "p": shared_values(fams, "n_phones"),
        "a": shared_values(fams, "n_addr"),
    }
    pairs, skipped_blocks = candidate_pairs(fams)
    print(f"{len(pairs)} candidate pairs ({skipped_blocks} oversized blocks skipped)")
    edges, cannot, stats = {}, set(), collections.Counter()
    for i, j in pairs:
        ev = score_pair(fams[i], fams[j], shared, stats)
        if ev and ev[0] is None:
            cannot.add((i, j))
        elif ev:
            edges[(i, j)] = ev
    print(f"{len(edges)} matching pairs, {len(cannot)} related/vetoed pairs kept apart")
    apart_cache = {}

    def apart(i, j):
        """Must records i and j stay in different groups?  Kept-apart pairs:
        yes; matched pairs: no; any other pair (linked only through third
        records) gets the Jr./Sr. and member-verdict checks on demand."""
        key = (i, j) if i < j else (j, i)
        if key in cannot:
            return True
        if key in edges:
            return False
        if key not in apart_cache:
            f, g = fams[key[0]], fams[key[1]]
            apart_cache[key] = suffix_conflict(f, g) or member_verdict(member_evidence(f, g)) is not None
        return apart_cache[key]

    groups, crossing_weak, blocked = build_groups(fams, edges, cannot, apart)
    grouped = {m for g in groups for m in [g["survivor"]] + [x for x, _, _ in g["members"]]}
    # Records with a tier 1-3 match that ended up in no group (every link
    # was blocked by a kept-apart pair): listed for manual review.
    orphans = {}
    for (i, j), (tier, reasons) in edges.items():
        if tier > 3:
            continue
        for a, b in ((i, j), (j, i)):
            if a not in grouped and (a not in orphans or tier < orphans[a][1]):
                orphans[a] = (b, tier, reasons)

    # Reason codes: one per distinct (tier, reasons) combination.
    reason_counts = collections.Counter((t, r) for g in groups for _, t, r in g["members"])
    reason_ids = {
        key: f"R{n:03d}"
        for n, key in enumerate(sorted(reason_counts, key=lambda k: (k[0], -reason_counts[k], k[1])), 1)
    }

    def sort_key(g):
        s = fams[g["survivor"]]
        return (g["tier"], text(s["last"]).lower(), text(s["first"]).lower(), s["duid"])

    groups.sort(key=sort_key)

    def fam_cols(f):
        return [
            f["duid"],
            text(f["last"]),
            text(f["first"]),
            text(f["mailing"]),
            text(f["org"]),
            f["org_id"],
            text(f["registered"]),
            text(f["group"]),
            fmt_date(f["membership"]),
            fmt_date(f["unregistered"]),
            fmt_date(f["last_updated"]),
            fmt_date(f["activity"]),
            text(f["addr1"]) + (", " + text(f["addr2"]) if text(f["addr2"]) else ""),
            text(f["city"]),
            text(f["state"]),
            text(f["zip"]),
            ", ".join(sorted({text(f["home_phone"]), text(f["primary_phone"])} - {""})),
            text(f["email"]),
            "; ".join(m["label"] for m in f["members"]),
        ]

    fam_hdr = [
        "DUID",
        "Last Name",
        "First Names",
        "Mailing Name",
        "Organization",
        "Org ID",
        "Registered",
        "Group",
        "Membership Date",
        "Date Unregistered",
        "Family Last Updated",
        "Activity Date Used",
        "Address",
        "City",
        "State",
        "ZIP",
        "Phones",
        "Email",
        "Members",
    ]
    fam_w = [9, 16, 18, 28, 30, 8, 10, 12, 12, 12, 12, 12, 30, 14, 6, 11, 26, 30, 60]
    assert len(fam_hdr) == len(fam_w)

    dup_rows, pair_rows = [], []
    types = collections.Counter()
    type_by_tier = collections.Counter()
    for gn, g in enumerate(groups, 1):
        s = fams[g["survivor"]]
        conf = CONFIDENCE[g["tier"]]
        dup_rows.append([gn, conf, "", "SURVIVE", s["duid"]] + fam_cols(s) + ["", ""])
        for m, tier, reasons in sorted(g["members"], key=lambda x: (x[1], fams[x[0]]["duid"])):
            f = fams[m]
            rid = reason_ids[(tier, reasons)]
            dup_rows.append([gn, conf, CONFIDENCE[tier], "merge", s["duid"]] + fam_cols(f) + [rid, "; ".join(reasons)])
            mt = match_type(s, f)
            types[mt] += 1
            type_by_tier[(mt, tier)] += 1
            pair_rows.append(
                [
                    s["duid"],
                    f["duid"],
                    CONFIDENCE[tier],
                    rid,
                    gn,
                    text(s["last"]),
                    text(s["first"]),
                    text(s["org"]),
                    text(f["last"]),
                    text(f["first"]),
                    text(f["org"]),
                    mt,
                ]
            )

    # Merge Pairs: by each pair's own confidence, then survivor last name.
    pair_rows.sort(
        key=lambda r: (r[2] != CONFIDENCE[1], r[2] != CONFIDENCE[2], r[2] != CONFIDENCE[3], r[5].lower(), r[4])
    )
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Duplicates"
    write_sheet(
        ws,
        ["Group", "Group Confidence", "Match Confidence", "Role", "Survivor DUID"]
        + fam_hdr
        + ["Reason Code", "Reasons"],
        dup_rows,
        [7, 14, 14, 9, 10] + fam_w + [9, 60],
    )
    write_sheet(
        wb.create_sheet("Merge Pairs"),
        [
            "Survivor DUID (1st)",
            "Merged DUID (2nd)",
            "Confidence",
            "Reason Code",
            "Group",
            "Survivor Last",
            "Survivor First",
            "Survivor Org",
            "Merged Last",
            "Merged First",
            "Merged Org",
            "Match Type",
        ],
        pair_rows,
        [12, 12, 14, 9, 7, 16, 18, 30, 16, 18, 30, 55],
    )
    group_of = {}
    for gn, g in enumerate(groups, 1):
        for m in [g["survivor"]] + [x for x, _, _ in g["members"]]:
            group_of[m] = gn
    orphan_rows = []
    for a, (b, tier, reasons) in sorted(orphans.items(), key=lambda kv: (kv[1][1], text(fams[kv[0]]["last"]).lower())):
        f, o = fams[a], fams[b]
        orphan_rows.append(
            [
                f["duid"],
                text(f["last"]),
                text(f["first"]),
                text(f["org"]),
                text(f["registered"]),
                text(f["group"]),
                o["duid"],
                group_of.get(b, ""),
                CONFIDENCE[tier],
                "; ".join(reasons),
                "linking it would put Families shown to be related or different people into one group",
            ]
        )
    reason_rows = [
        [rid, CONFIDENCE[tier], "; ".join(reasons), reason_counts[(tier, reasons)]]
        for (tier, reasons), rid in sorted(reason_ids.items(), key=lambda kv: kv[1])
    ]
    write_sheet(
        wb.create_sheet("Reasons"),
        ["Reason Code", "Confidence", "Reasons", "Families Merged"],
        reason_rows,
        [9, 14, 90, 10],
    )
    write_sheet(
        wb.create_sheet("Not Grouped"),
        [
            "DUID",
            "Last Name",
            "First Names",
            "Organization",
            "Registered",
            "Group",
            "Best Match DUID",
            "Best Match Group",
            "Match Confidence",
            "Reasons",
            "Why Not Grouped",
        ],
        orphan_rows,
        [9, 16, 18, 30, 10, 12, 12, 10, 14, 60, 60],
    )

    # Summary sheet + Markdown.
    tier_fams = collections.Counter(t for g in groups for _, t, _ in g["members"])
    tier_groups = collections.Counter(g["tier"] for g in groups)
    sizes = collections.Counter(len(g["members"]) + 1 for g in groups)
    summ = wb.create_sheet("Summary")
    lines = [
        ["Input families", len(fams)],
        ["Input members (after dropping surname-as-first-name placeholder entries)", n_members],
        ["Families with no usable members (none exported, or only placeholders)", no_members],
        ["Families with usable enhanced-export fields", with_enh],
        ["Placeholder records excluded (Loose Collection, Unknown Family, ...)", sum(f["placeholder"] for f in fams)],
        ["Catch-all Families (members span 6+ surnames; member evidence ignored)", sum(f["catchall"] for f in fams)],
        ["Candidate pairs examined", len(pairs)],
        ["Matching pairs", len(edges)],
        ["Duplicate groups", len(groups)],
        ["Families to merge away", sum(tier_fams.values())],
        ["Low (review) links not used because one or both records are already in a stronger group", crossing_weak],
        [
            "Very high/High/Medium links not used because they would group two Families shown to be "
            "related or different people",
            blocked,
        ],
        ["Records with a match but in no group (see the Not Grouped sheet)", len(orphan_rows)],
        [
            "Bulk days ignored for 'Date Unregistered' (Archdiocese-wide)",
            ", ".join(str(d) for d in sorted(k for k in unreg_bulk if not isinstance(k, tuple))),
        ],
        ["Bulk parish-days ignored for 'Date Unregistered'", sum(isinstance(k, tuple) for k in unreg_bulk)],
        ["Bulk days + parish-days ignored for 'Date Last Updated'", len(upd_bulk)],
        [],
        ["Pairs rejected or reclassified by member evidence", "Pairs"],
    ]
    lines += [[k, v] for k, v in stats.most_common()]
    lines += [[], ["Confidence", "Groups (by weakest member)", "Families merged (by own match)"]]
    lines += [[CONFIDENCE[t], tier_groups[t], tier_fams[t]] for t in sorted(CONFIDENCE)]
    lines += [[], ["Match type", "Families merged"] + [CONFIDENCE[t] for t in sorted(CONFIDENCE)]]
    lines += [[mt, n] + [type_by_tier[(mt, t)] for t in sorted(CONFIDENCE)] for mt, n in types.most_common()]
    lines += [[], ["Group size (records)", "Groups"]] + [[k, v] for k, v in sorted(sizes.items())]
    for ln in lines:
        summ.append(ln)
    summ.column_dimensions["A"].width = 95
    for col in "BCDEF":
        summ.column_dimensions[col].width = 16

    xlsx = out_dir / "archlou-duplicate-families.xlsx"
    wb.save(xlsx)
    md = out_dir / "summary.md"
    md.write_text("\n".join("| " + " | ".join(str(c) for c in ln) + " |" if ln else "" for ln in lines) + "\n")
    print(f"wrote {xlsx} and {md}")
    for ln in lines:
        print("  ", *ln)


if __name__ == "__main__":
    main()
