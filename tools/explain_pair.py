#!/usr/bin/env python3
"""Explain how dedup.py treats two Family DUIDs.

Prints the pair's score (tier and reasons, a kept-apart verdict, or no
match), the member evidence behind it, each Family's members, and -- when
a report workbook is given -- which group each DUID landed in.

Usage:
  tools/explain_pair.py FAMILIES.xlsx MEMBERS.xlsx DUID_A DUID_B \
      --cache-dir DATA_DIR/output [--enhanced ENHANCED.xlsx] [--report REPORT.xlsx]

Use the same --cache-dir as the dedup.py run (its --out-dir) so the cached
exports are reused.  The output contains personal data: read it in the
terminal, don't paste it into issues, commits or docs.
"""

import argparse
import collections
import importlib.util
from pathlib import Path

import openpyxl


def load_dedup():
    """Import dedup.py from the repository root (no sys.path changes)."""
    path = Path(__file__).resolve().parent.parent / "dedup.py"
    spec = importlib.util.spec_from_file_location("dedup", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def report_groups(path):
    """Map DUID -> (group number, role, match confidence, reasons) from a report workbook."""
    wb = openpyxl.load_workbook(path, read_only=True)
    rows = wb["Duplicates"].iter_rows(values_only=True)
    hdr = {h: i for i, h in enumerate(next(rows))}
    out = {}
    for r in rows:
        out[r[hdr["DUID"]]] = (r[0], r[hdr["Role"]], r[hdr["Match Confidence"]], r[hdr["Reasons"]])
    not_grouped = {r[0] for r in list(wb["Not Grouped"].iter_rows(values_only=True))[1:]}
    return out, not_grouped


def main():
    """Load the exports, score the pair, and print the explanation."""
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("families")
    ap.add_argument("members")
    ap.add_argument("duid_a", type=int)
    ap.add_argument("duid_b", type=int)
    ap.add_argument("--cache-dir", required=True, help="the --out-dir used for dedup.py (holds its caches)")
    ap.add_argument("--enhanced")
    ap.add_argument("--report", help="report workbook from dedup.py, to show group membership")
    args = ap.parse_args()

    d = load_dedup()
    fams = d.load(Path(args.families), args.cache_dir)
    members = d.load_members(Path(args.members), args.cache_dir)
    enhanced = d.load_enhanced(Path(args.enhanced), args.cache_dir) if args.enhanced else {}
    d.enrich(fams, members, enhanced)
    shared = {
        "e": d.shared_values(fams, "n_emails"),
        "p": d.shared_values(fams, "n_phones"),
        "a": d.shared_values(fams, "n_addr"),
    }
    by_duid = {f["duid"]: f for f in fams}
    f, g = by_duid[args.duid_a], by_duid[args.duid_b]

    result = d.score_pair(f, g, shared, collections.Counter())
    if result is None:
        print("score: no match")
    elif result[0] is None:
        print(f"score: KEPT APART -- {result[1]}")
    else:
        print(f"score: {d.CONFIDENCE[result[0]]} -- {'; '.join(result[1])}")
    ev = d.member_evidence(f, g)
    print("member evidence:", {k: v for k, v in ev.items() if v and not k.endswith(("_list_f", "_list_g"))})
    for x in (f, g):
        print(
            f"\n{x['duid']}: {x['first']} {x['last']} | {x['org']} | registered={x['registered']} "
            f"group={x['group']} | activity={d.fmt_date(x['activity'])}"
        )
        for m in x["members"]:
            print(f"    {m['label']}  role={m['role']}{' (guessed)' if m['guessed'] else ''}")
    if args.report:
        groups, not_grouped = report_groups(args.report)
        print()
        for duid in (args.duid_a, args.duid_b):
            where = groups.get(duid) or ("Not Grouped sheet" if duid in not_grouped else "not in the report")
            print(f"{duid}: {where}")


if __name__ == "__main__":
    main()
